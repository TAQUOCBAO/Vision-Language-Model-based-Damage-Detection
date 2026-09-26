#!/usr/bin/env python3
"""Build a 16:9 widescreen PowerPoint presentation for Structural Damage VLM Diagnosis."""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

os.makedirs('ppt', exist_ok=True)

# -------------------------------------------------------------
# Color Palette (Dark Tech Palette)
# -------------------------------------------------------------
BG_COLOR = RGBColor(15, 23, 42)        # #0F172A Slate 900
CARD_BG = RGBColor(30, 41, 59)        # #1E293B Slate 800
TEXT_MAIN = RGBColor(248, 250, 252)   # #F8FAFC
TEXT_MUTED = RGBColor(148, 163, 184)  # #94A3B8
ACCENT_CYAN = RGBColor(56, 189, 248)  # #38BDF8 Sky Blue
ACCENT_GREEN = RGBColor(52, 211, 153) # #34D399 Emerald
ACCENT_AMBER = RGBColor(251, 191, 36) # #FBBF24 Amber
ACCENT_PURPLE = RGBColor(168, 85, 247)# #A855F7 Purple
BORDER_COLOR = RGBColor(51, 65, 85)   # #334155

def create_presentation():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Blank slide layout

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="Project Presentation"):
        # Header box
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.9))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        # Category / Pill
        p1 = tf.paragraphs[0]
        p1.text = f"• {category_text.upper()} •"
        p1.font.size = Pt(10)
        p1.font.bold = True
        p1.font.color.rgb = ACCENT_CYAN
        
        # Title
        p2 = tf.add_paragraph()
        p2.text = title_text
        p2.font.size = Pt(22)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_MAIN
        p2.space_before = Pt(3)

    # =========================================================
    # SLIDE 1: Title Slide
    # =========================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Accent Card on Title
    card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.2), Inches(10.9), Inches(5.1))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = BORDER_COLOR
    card.line.width = Pt(1.5)

    tb = s1.shapes.add_textbox(Inches(1.8), Inches(1.6), Inches(9.7), Inches(4.3))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "MULTIMODAL AI & COMPUTER VISION FOR CIVIL INFRASTRUCTURE"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    p = tf.add_paragraph()
    p.text = "Structural Damage Image-Text Diagnosis"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = TEXT_MAIN
    p.space_before = Pt(8)

    p = tf.add_paragraph()
    p.text = "Fine-Tuning Qwen3-VL-8B-Instruct via LoRA with 3-Tier Data Bake-Off and Low-VRAM Deployment"
    p.font.size = Pt(16)
    p.font.color.rgb = ACCENT_GREEN
    p.space_before = Pt(8)

    p = tf.add_paragraph()
    p.text = "• Model: Qwen3-VL-8B-Instruct  |  • Holdout Multilabel F1: 0.9528  |  • Framework: LLaMA-Factory + LoRA (BF16)\n• Published LoRA: nhantran214/damage-vlm-qwen3vl-8b-lora  |  • Target GPU: RTX 5880 Ada (49GB)"
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(28)

    # =========================================================
    # SLIDE 2: Executive Summary & Objective
    # =========================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "1. Executive Summary & Mission", "Overview")

    # Left Column (Overview & JSON output)
    c1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.4))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = BORDER_COLOR

    tb = s2.shapes.add_textbox(Inches(1.1), Inches(1.7), Inches(5.1), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Mission Objective"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    p = tf.add_paragraph()
    p.text = "Build an offline AI pipeline that takes structural inspection photos and automatically generates forensic diagnosis reports formatted strictly as valid JSON with zero conversational filler."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MAIN
    p.space_before = Pt(8)

    p = tf.add_paragraph()
    p.text = "Standard Output Schema:"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_AMBER
    p.space_before = Pt(12)

    p = tf.add_paragraph()
    p.text = "{\n  \"image_id\": \"00002\",\n  \"damage_categories\": [\"cracks\", \"spalling\"],\n  \"description\": \"Longitudinal fracture along beam with adjacent surface spalling and exposed matrix.\"\n}"
    p.font.size = Pt(11)
    p.font.name = "Courier New"
    p.font.color.rgb = RGBColor(226, 232, 240)
    p.space_before = Pt(6)

    # Right Column (Key Results)
    c2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.4))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = BORDER_COLOR

    tb = s2.shapes.add_textbox(Inches(7.1), Inches(1.7), Inches(5.1), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Core Quantitative Achievements"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    items = [
        ("0.9528 Multilabel F1 Score", "Holdout validation (n=237) across 10 closed damage classes with zero leakage."),
        ("0.8268 Winning Weighted Score", "Formula: 0.6 * F1 + 0.4 * METEOR. Won over all augmentation variants."),
        ("Frozen Vision + Projector", "Preserves pretrained visual representations and saves ~18GB VRAM."),
        ("Dual Deployment Modes", "Full BF16 high-precision mode (~18GB) + 4-Bit low-VRAM mode (~5.8GB for RTX 3060).")
    ]
    for title, desc in items:
        p = tf.add_paragraph()
        p.text = f"• {title}: "
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = ACCENT_CYAN
        p.space_before = Pt(10)
        
        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MUTED

    # =========================================================
    # SLIDE 3: Domain Complexity & 10-Class Taxonomy
    # =========================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "2. Problem Framing & 10-Class Fixed Taxonomy", "Domain Challenge")

    # Left: Domain Challenges
    c1 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.4))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = BORDER_COLOR

    tb = s3.shapes.add_textbox(Inches(1.1), Inches(1.7), Inches(5.1), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Domain-Specific Challenges"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_AMBER

    challenges = [
        ("Multi-Defect Co-occurrence", ">64% of inspection photos exhibit 2+ overlapping damage types (e.g. spalling + rebar corrosion). Single-label classification fails completely."),
        ("Spatial Orientation Sensitivity", "Captions contain critical directional terms ('vertical crack', 'left-to-right'). Random image flipping invalidates ground-truth text semantics!"),
        ("Long-Tailed Rare Classes", "Small dataset (~1,200 images) with severe imbalance on rare classes like honeycombing, efflorescence, and potholes.")
    ]
    for title, desc in challenges:
        p = tf.add_paragraph()
        p.text = f"• {title}"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = RGBColor(239, 68, 68)
        p.space_before = Pt(12)

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(2)

    # Right: Taxonomy Table
    c2 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.4))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = BORDER_COLOR

    tb = s3.shapes.add_textbox(Inches(7.1), Inches(1.7), Inches(5.1), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "10 Closed Damage Classes (configs/labels_v1.yaml)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    p = tf.add_paragraph()
    p.text = "1. spalling (Bong tróc bê tông)\n2. cracks (Vết nứt kết cấu)\n3. corrosion (Ăn mòn / Gỉ sét kim loại)\n4. voids (Lỗ rỗng / Hốc vật liệu)\n5. exposed_rebar (Lộ cốt thép)\n6. peeling (Bong tróc lớp phủ / Vảy sơn)\n7. potholes (Ổ gà / Khuyết tật mặt đường)\n8. honeycomb (Rỗ tổ ong bê tông)\n9. looseness (Lỏng lẻo / Tách lớp)\n10. efflorescence (Phấn hóa / Muối kiềm đóng vôi)"
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MAIN
    p.space_before = Pt(10)

    # =========================================================
    # SLIDE 4: Qwen3-VL Architecture & Adaptation (KEY SLIDE WITH IMAGE)
    # =========================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "3. Qwen3-VL-8B Architecture & LoRA Adaptation Deep-Dive", "Model Architecture")

    # Insert Architecture Diagram
    diag_path = 'ppt/figures/qwen3_vl_architecture.png'
    if os.path.exists(diag_path):
        s4.shapes.add_picture(diag_path, Inches(0.8), Inches(1.5), width=Inches(11.7))

    # =========================================================
    # SLIDE 5: Architectural Breakdown & Key Decisions
    # =========================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "4. Architectural Breakdown & VRAM-Safe Strategy", "Architecture Details")

    grid_cards = [
        ("Vision Tower (ViT)", "Native Dynamic Resolution ViT with 2D Rotary Position Embeddings (RoPE). Splits images into variable aspect-ratio patches.\n\nStrategy: 100% FROZEN (0 parameters trained). Preserves high-quality vision representations and saves ~18 GB VRAM."),
        ("Spatial Projector", "Compresses 2x2 vision patch tokens into 1 unified token before feeding into the language model. Projects visual embeddings to 4096 dimensions.\n\nStrategy: 100% FROZEN. Avoids catastrophic forgetting of vision-language alignment."),
        ("Qwen3 LLM Backbone", "36 Transformer Decoder layers with Grouped Query Attention (GQA), SwiGLU activations, and RMSNorm (8B parameters).\n\nLoRA Adaptation: Injected into q_proj, k_proj, v_proj, o_proj, gate_proj, up_proj, down_proj (Rank 64, Alpha 128)."),
        ("Autoregressive Output", "Outputs structured forensic JSON string directly. Prompt enforces single JSON object with closed-class set.\n\nDownstream Alignment: Followed by deterministic post-processing (JSON repair, synonym mapping, and text-category synchronization).")
    ]

    for idx, (title, content) in enumerate(grid_cards):
        row = idx // 2
        col = idx % 2
        left = Inches(0.8 + col * 6.0)
        top = Inches(1.5 + row * 2.8)

        c = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.7), Inches(2.6))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = BORDER_COLOR

        tb = s5.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), Inches(5.3), Inches(2.3))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"{idx+1}. {title}"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = ACCENT_CYAN

        p = tf.add_paragraph()
        p.text = content
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(4)

    # =========================================================
    # SLIDE 6: Data Engineering & Silver Labeling
    # =========================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "5. Data Engineering & Stratified Multi-Label Split", "Data Engineering")

    c1 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.4))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = BORDER_COLOR

    tb = s6.shapes.add_textbox(Inches(1.1), Inches(1.7), Inches(5.1), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Silver Label Synthesis Rule"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    p = tf.add_paragraph()
    p.text = "Because raw inspection files only contained conversational VQA records, silver targets are built automatically by taking the union:"
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MAIN
    p.space_before = Pt(6)

    p = tf.add_paragraph()
    p.text = "L_i = Prefix_Seeds(filename) ∪ Text_Keywords(description)"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.name = "Courier New"
    p.font.color.rgb = ACCENT_AMBER
    p.space_before = Pt(8)

    p = tf.add_paragraph()
    p.text = "• Prefix Seeds: Letter prefixes (crack_, gangf_, xiu_) identify the original defect batch.\n• Longest-First Matching: Prioritizes longer phrases (e.g. 'exposed rebar' before 'exposed') to avoid false alarms."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(8)

    # Right: Stratified Split
    c2 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.4))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = BORDER_COLOR

    tb = s6.shapes.add_textbox(Inches(7.1), Inches(1.7), Inches(5.1), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Frozen Multi-Label Stratified Split"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    p = tf.add_paragraph()
    p.text = "• Algorithm: MultilabelStratifiedShuffleSplit (80/20, seed 42).\n• Train Set: 963 images | Validation Set: 237 images.\n• Artifact: data/processed/splits/v1.json.\n• Zero Leakage Guarantee: Split is frozen once and strictly reused for all tier bake-offs. Validation images never enter augmentation or paraphrase pipelines."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MAIN
    p.space_before = Pt(8)

    # =========================================================
    # SLIDE 7: Three-Tier Data Bake-Off (A/B/C)
    # =========================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "6. Three-Tier Data Augmentation Bake-Off", "Experimental Design")

    tiers = [
        ("Tier A: Clean Baseline SFT", "WINNER", ACCENT_GREEN, 
         "• 963 clean train pairs.\n• Zero synthetic image distortion.\n• Raw longest description per image.\n• Lowest label noise and highest category precision."),
        ("Tier B: Orientation-Safe Aug", "RUNNER-UP", ACCENT_CYAN, 
         "• 1,926 pairs (Original + Aug).\n• Photometric: Brightness, Contrast, Blur, Noise.\n• Geometric Flip Rule: Strictly forbidden if description mentions 'vertical', 'horizontal', or 'left-to-right'."),
        ("Tier C: Paraphrase & External", "RANK 3", ACCENT_AMBER, 
         "• 2,889+ expanded pairs.\n• 3x Local text paraphrasing.\n• External rare-class image injection.\n• Higher linguistic diversity but synthetic noise hurt classification F1.")
    ]

    for idx, (title, badge, color, details) in enumerate(tiers):
        left = Inches(0.8 + idx * 4.0)
        c = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.5), Inches(3.7), Inches(5.4))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = color
        c.line.width = Pt(1.5)

        tb = s7.shapes.add_textbox(left + Inches(0.2), Inches(1.7), Inches(3.3), Inches(5.0))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"[{badge}]"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = color

        p = tf.add_paragraph()
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(4)

        p = tf.add_paragraph()
        p.text = details
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(12)

    # =========================================================
    # SLIDE 8: Experimental Results & Winner Selection
    # =========================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "7. Holdout Evaluation & Winner Selection Results", "Experimental Results")

    # Left: Comparison Chart
    chart_path = 'ppt/figures/bakeoff_comparison.png'
    if os.path.exists(chart_path):
        s8.shapes.add_picture(chart_path, Inches(0.8), Inches(1.5), width=Inches(6.6))

    # Right: Results Table & Formula
    c2 = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.6), Inches(1.5), Inches(4.9), Inches(5.4))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = BORDER_COLOR

    tb = s8.shapes.add_textbox(Inches(7.8), Inches(1.7), Inches(4.5), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Holdout Scores (N=237)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    p = tf.add_paragraph()
    p.text = "• Tier A (Winner):\n  F1: 0.9528 | METEOR: 0.6379 | Score: 0.82684\n• Tier B (Runner-Up):\n  F1: 0.9495 | METEOR: 0.6427 | Score: 0.82681\n• Tier C (Rank 3):\n  F1: 0.9378 | METEOR: 0.6136 | Score: 0.80814"
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_MAIN
    p.space_before = Pt(8)

    p = tf.add_paragraph()
    p.text = "Winner Selection Formula:"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_AMBER
    p.space_before = Pt(12)

    p = tf.add_paragraph()
    p.text = "Score = 0.6 * F1 + 0.4 * METEOR\nSafety Gate: F1 >= 0.97 * Best_F1 (Passed)"
    p.font.size = Pt(11)
    p.font.name = "Courier New"
    p.font.color.rgb = ACCENT_GREEN
    p.space_before = Pt(4)

    # =========================================================
    # SLIDE 9: Deterministic Post-Processing & Alignment
    # =========================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "8. Deterministic Post-Processing & Alignment Stack", "Post-Processing")

    steps = [
        ("1. JSON Syntax Repair", "Extracts JSON from markdown fences (```json) or outer curly braces { ... }, tolerating conversational preambles."),
        ("2. Synonym Canonicalization", "Normalizes free-form predicted damage terms back to the exact 10 closed-class set (e.g. 'rusty steel' -> 'corrosion')."),
        ("3. Category-Text Alignment", "Synchronizes categories with description: automatically appends standard morphological sentences if a predicted defect is omitted from text."),
        ("4. Sentence Deduplication", "Filters out duplicate phrases and normalizes whitespace while preserving all technical dimensions and measurements.")
    ]

    for idx, (title, desc) in enumerate(steps):
        row = idx // 2
        col = idx % 2
        left = Inches(0.8 + col * 6.0)
        top = Inches(1.5 + row * 2.8)

        c = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.7), Inches(2.6))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = BORDER_COLOR

        tb = s9.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), Inches(5.3), Inches(2.3))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = ACCENT_CYAN

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(6)

    # =========================================================
    # SLIDE 10: Production Retraining on data_v2 (1,200 Images)
    # =========================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "9. Production Retrain on Updated Corpus (data_v2)", "Production Scale")

    c1 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.4))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = BORDER_COLOR

    tb = s10.shapes.add_textbox(Inches(1.1), Inches(1.7), Inches(5.1), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Full Retrain Execution"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    p = tf.add_paragraph()
    p.text = "• Dataset Source: data_v2/dataset (1,200 total images).\n• Recipe: Winning Tier A SFT (clean multi-label conversations).\n• Training Length: 3.0 Epochs (225 Optimization Steps).\n• Convergence: Training loss converged smoothly to 0.249.\n• Output Adapter: outputs/runs/full_winner_v2.\n• Final Submission: outputs/submissions/submission.json."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MAIN
    p.space_before = Pt(10)

    # Right: Loss Progression Stats
    c2 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.4))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = BORDER_COLOR

    tb = s10.shapes.add_textbox(Inches(7.1), Inches(1.7), Inches(5.1), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Training Loss Convergence Trajectory"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    p = tf.add_paragraph()
    p.text = "• Step 25 (Epoch 0.3): Loss 1.450\n• Step 50 (Epoch 0.7): Loss 0.780\n• Step 100 (Epoch 1.3): Loss 0.420\n• Step 150 (Epoch 2.0): Loss 0.310\n• Step 200 (Epoch 2.7): Loss 0.260\n• Step 225 (Epoch 3.0): Final Loss 0.249"
    p.font.size = Pt(12)
    p.font.name = "Courier New"
    p.font.color.rgb = RGBColor(226, 232, 240)
    p.space_before = Pt(10)

    # =========================================================
    # SLIDE 11: Deployment & Low-VRAM Execution
    # =========================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11)
    add_header(s11, "10. Low-VRAM Deployment & Hugging Face Hub", "Deployment")

    # Left: VRAM Chart
    vram_path = 'ppt/figures/vram_comparison.png'
    if os.path.exists(vram_path):
        s11.shapes.add_picture(vram_path, Inches(0.8), Inches(1.5), width=Inches(6.0))

    # Right: Hugging Face Info
    c2 = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.0), Inches(1.5), Inches(5.5), Inches(5.4))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = BORDER_COLOR

    tb = s11.shapes.add_textbox(Inches(7.2), Inches(1.7), Inches(5.1), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Hugging Face Model Publication"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    p = tf.add_paragraph()
    p.text = "Repo ID: nhantran214/damage-vlm-qwen3vl-8b-lora\nSize: ~0.7 GiB (LoRA Adapter Weights only)"
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_MAIN
    p.space_before = Pt(6)

    p = tf.add_paragraph()
    p.text = "Download Command:"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ACCENT_AMBER
    p.space_before = Pt(10)

    p = tf.add_paragraph()
    p.text = "huggingface-cli download \\\n  nhantran214/damage-vlm-qwen3vl-8b-lora \\\n  --local-dir ./models/damage-vlm-qwen3vl-8b-lora"
    p.font.size = Pt(10)
    p.font.name = "Courier New"
    p.font.color.rgb = RGBColor(226, 232, 240)
    p.space_before = Pt(4)

    p = tf.add_paragraph()
    p.text = "Low-VRAM 4-Bit Inference Command:"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p.space_before = Pt(10)

    p = tf.add_paragraph()
    p.text = "python scripts/08_infer_full_winner.py --load-in-4bit \\\n  --adapter ./models/damage-vlm-qwen3vl-8b-lora \\\n  --image-dir data_v2/dataset/image"
    p.font.size = Pt(10)
    p.font.name = "Courier New"
    p.font.color.rgb = RGBColor(226, 232, 240)
    p.space_before = Pt(4)

    # =========================================================
    # SLIDE 12: Summary & Conclusion
    # =========================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12)
    add_header(s12, "11. Summary, Key Takeaways & Conclusion", "Summary")

    takeaways = [
        ("Methodological Rigor", "Empirically demonstrated that high-quality clean SFT (Tier A) surpasses noisy synthetic data augmentation for multi-defect engineering diagnosis."),
        ("Engineering Reliability", "4-stage post-processing stack guarantees 100% compliant JSON schema output and enforces bidirectional category-description consistency."),
        ("High Accessibility", "Dual deployment options allow uncompromised BF16 execution on server GPUs alongside 4-bit 5.8GB VRAM execution on consumer cards.")
    ]

    for idx, (title, text) in enumerate(takeaways):
        left = Inches(0.8 + idx * 4.0)
        c = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.5), Inches(3.7), Inches(5.4))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = BORDER_COLOR

        tb = s12.shapes.add_textbox(left + Inches(0.2), Inches(1.7), Inches(3.3), Inches(5.0))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"{idx+1}. {title}"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = ACCENT_CYAN

        p = tf.add_paragraph()
        p.text = text
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(12)

    # Save presentation
    out_file = 'ppt/Structural_Damage_VLM_Diagnosis.pptx'
    prs.save(out_file)
    print(f"Successfully generated PowerPoint file: {out_file}")

if __name__ == '__main__':
    create_presentation()
